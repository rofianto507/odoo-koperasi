"""Impor anggota dari Excel: unduh templat -> isi -> unggah -> periksa (pratinjau) -> impor.

Anggota yang diimpor dibuat seperti lewat form (berstatus Draft, kontak dan riwayat gaji ikut
terbentuk), dengan hak pengguna yang mengimpor: admin mitra hanya untuk mitranya sendiri.
"""
import base64
import io
import json
import re
from datetime import date, datetime

from odoo import api, fields, models, _
from odoo.exceptions import UserError

from ..models.koperasi_anggota import BANK_SELECTION

BARIS_MAKS = 2000

# (field, judul kolom, wajib, jenis, keterangan di lembar Petunjuk)
KOLOM = [
    ('name', 'Nama Lengkap', True, 'teks', 'Nama sesuai KTP.'),
    ('nip', 'NIP / ID Karyawan', True, 'kode', 'Nomor pegawai di perusahaan. Tidak boleh sama antar pegawai.'),
    ('gaji_bersih', 'Gaji Bersih', True, 'uang', 'Gaji bersih per bulan dalam rupiah, angka saja. Dipakai menghitung plafon pinjaman.'),
    ('nik', 'NIK (KTP)', False, 'kode', '16 digit.'),
    ('tempat_lahir', 'Tempat Lahir', False, 'teks', ''),
    ('tanggal_lahir', 'Tanggal Lahir', False, 'tanggal', 'Format tanggal Excel, atau teks hh/bb/tttt (mis. 17/08/1990).'),
    ('jenis_kelamin', 'Jenis Kelamin', False, 'pilihan', 'Laki-laki atau Perempuan (boleh L / P).'),
    ('telepon', 'Telp / WA', False, 'kode', 'Mis. 081234567890.'),
    ('email', 'Email', False, 'teks', ''),
    ('jabatan', 'Jabatan', False, 'teks', ''),
    ('departemen', 'Departemen / Unit', False, 'teks', ''),
    ('status_pegawai', 'Status Pegawai', False, 'pilihan', 'Kontrak atau Tetap.'),
    ('tanggal_masuk_kerja', 'Mulai Bekerja', False, 'tanggal', 'Format tanggal Excel, atau teks hh/bb/tttt.'),
    ('bank_tujuan', 'Bank', False, 'pilihan', 'Pilih dari daftar. Bank di luar daftar boleh diketik; akan disimpan sebagai Bank Lainnya.'),
    ('nomor_rekening', 'Nomor Rekening', False, 'kode', 'Angka saja (spasi dan tanda hubung diperbolehkan).'),
    ('atas_nama', 'Atas Nama Rekening', False, 'teks', 'Bila kosong, diisi nama anggota.'),
    ('simsula', 'Simpanan Sukarela / Bulan', False, 'uang', 'Rupiah per bulan, boleh kosong.'),
    ('street', 'Alamat', False, 'teks', ''),
    ('city', 'Kota', False, 'teks', ''),
]
PILIHAN = {
    'jenis_kelamin': [('l', 'Laki-laki'), ('p', 'Perempuan')],
    'status_pegawai': [('kontrak', 'Kontrak'), ('tetap', 'Tetap')],
    'bank_tujuan': [b for b in BANK_SELECTION if b[0] != 'lainnya'],
}
SINONIM = {
    'jenis_kelamin': {'l': 'l', 'lakilaki': 'l', 'pria': 'l', 'lk': 'l', 'p': 'p', 'perempuan': 'p', 'wanita': 'p', 'pr': 'p'},
    'status_pegawai': {'kontrak': 'kontrak', 'pkwt': 'kontrak', 'tetap': 'tetap', 'pkwtt': 'tetap', 'permanen': 'tetap'},
}
LEMBAR_DATA = 'Anggota'


def _kunci(teks):
    """Bentuk ringkas untuk mencocokkan judul kolom dan pilihan: huruf kecil tanpa tanda baca."""
    return re.sub(r'[^a-z0-9]', '', str(teks or '').lower())


class KoperasiImporAnggotaWizard(models.TransientModel):
    _name = 'koperasi.impor.anggota.wizard'
    _description = 'Impor Anggota dari Excel'

    def _default_mitra(self):
        mitra = self.env.user.koperasi_mitra_id
        if mitra:
            return mitra
        aktif = self.env.context.get('default_mitra_id')
        return self.env['koperasi.mitra'].browse(aktif) if aktif else self.env['koperasi.mitra']

    state = fields.Selection(
        [('unggah', 'Unggah'), ('periksa', 'Periksa'), ('selesai', 'Selesai')], default='unggah', required=True)
    mitra_id = fields.Many2one(
        'koperasi.mitra', string='Mitra', required=True, default=_default_mitra,
        help='Semua anggota di berkas ini didaftarkan ke mitra ini.')
    mitra_terkunci = fields.Boolean(compute='_compute_mitra_terkunci')
    berkas = fields.Binary('Berkas Excel', attachment=False)
    nama_berkas = fields.Char()
    ajukan = fields.Boolean(
        'Langsung ajukan validasi', default=False,
        help='Anggota yang diimpor langsung berstatus Menunggu Validasi. Bila tidak dicentang, '
             'mereka berstatus Draft dan masih bisa diperiksa dulu.')
    line_ids = fields.One2many('koperasi.impor.anggota.line', 'wizard_id', string='Hasil Pemeriksaan')
    jumlah_ok = fields.Integer('Siap Diimpor', compute='_compute_jumlah')
    jumlah_galat = fields.Integer('Bermasalah', compute='_compute_jumlah')
    jumlah_impor = fields.Integer('Berhasil Diimpor', readonly=True)
    anggota_ids = fields.Many2many('koperasi.anggota', string='Anggota Diimpor', readonly=True)

    def _compute_mitra_terkunci(self):
        # admin mitra hanya mengimpor untuk mitranya sendiri
        self.mitra_terkunci = not self.env.user.has_group('koperasi.group_koperasi_operator')

    @api.depends('line_ids.status')
    def _compute_jumlah(self):
        for r in self:
            r.jumlah_ok = len(r.line_ids.filtered(lambda l: l.status == 'ok'))
            r.jumlah_galat = len(r.line_ids) - r.jumlah_ok

    def _ulang(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window', 'name': _('Impor Anggota'), 'res_model': self._name,
            'res_id': self.id, 'view_mode': 'form', 'target': 'new',
        }

    # ------------------------------------------------------------------
    # Templat
    # ------------------------------------------------------------------
    def action_unduh_templat(self):
        return {'type': 'ir.actions.act_url', 'url': '/koperasi/anggota/templat.xlsx', 'target': 'self'}

    @api.model
    def _buat_templat(self):
        import xlsxwriter  # noqa: PLC0415
        keluaran = io.BytesIO()
        wb = xlsxwriter.Workbook(keluaran, {'in_memory': True})
        ws = wb.add_worksheet(LEMBAR_DATA)
        wajib = wb.add_format({'bold': True, 'bg_color': '#71639E', 'font_color': '#FFFFFF', 'border': 1, 'text_wrap': True, 'valign': 'vcenter'})
        biasa = wb.add_format({'bold': True, 'bg_color': '#E7E6ED', 'border': 1, 'text_wrap': True, 'valign': 'vcenter'})
        format_kolom = {
            'kode': wb.add_format({'num_format': '@'}),          # teks: NIK/telepon tidak berubah jadi angka
            'uang': wb.add_format({'num_format': '#,##0'}),
            'tanggal': wb.add_format({'num_format': 'dd/mm/yyyy'}),
        }
        ws.set_row(0, 32)
        for i, (field, judul, harus, jenis, _ket) in enumerate(KOLOM):
            ws.write(0, i, judul + (' *' if harus else ''), wajib if harus else biasa)
            ws.set_column(i, i, max(16, len(judul) + 4), format_kolom.get(jenis))
            if field in PILIHAN:
                ws.data_validation(1, i, BARIS_MAKS, i, {
                    'validate': 'list', 'source': [label for _k, label in PILIHAN[field]],
                    # bank di luar daftar boleh diketik
                    'error_type': 'information' if field == 'bank_tujuan' else 'stop',
                    'error_title': 'Pilihan', 'error_message': 'Pilih salah satu nilai dari daftar.'})
            elif jenis == 'tanggal':
                ws.data_validation(1, i, BARIS_MAKS, i, {
                    'validate': 'date', 'criteria': 'between', 'minimum': date(1930, 1, 1), 'maximum': date(2100, 12, 31),
                    'error_type': 'warning', 'error_title': 'Tanggal', 'error_message': 'Isi dengan tanggal, mis. 17/08/1990.'})
        ws.freeze_panes(1, 2)

        p = wb.add_worksheet('Petunjuk')
        judul = wb.add_format({'bold': True, 'font_size': 14})
        tebal = wb.add_format({'bold': True, 'bg_color': '#E7E6ED', 'border': 1})
        sel = wb.add_format({'border': 1, 'text_wrap': True, 'valign': 'top'})
        p.set_column(0, 0, 28)
        p.set_column(1, 1, 10)
        p.set_column(2, 2, 78)
        p.write(0, 0, 'Petunjuk Pengisian Templat Anggota', judul)
        catatan = [
            'Isi data pada lembar "%s", satu pegawai per baris, mulai baris ke-2. Jangan mengubah judul kolom.' % LEMBAR_DATA,
            'Kolom bertanda * wajib diisi. Kolom lain boleh dikosongkan dan dilengkapi nanti.',
            'Setelah diunggah, sistem memeriksa tiap baris. Baris bermasalah tidak diimpor dan ditampilkan beserta alasannya.',
            'Anggota hasil impor berstatus Draft sampai diajukan dan disetujui operator koperasi.',
        ]
        for i, teks in enumerate(catatan):
            p.write(2 + i, 0, '%d.' % (i + 1))
            p.merge_range(2 + i, 1, 2 + i, 2, teks)
        baris = 3 + len(catatan)
        for i, h in enumerate(('Kolom', 'Wajib', 'Keterangan')):
            p.write(baris, i, h, tebal)
        for field, nama, harus, _jenis, ket in KOLOM:
            baris += 1
            if field in PILIHAN and field != 'bank_tujuan':
                ket = (ket + ' ' if ket else '') + 'Pilihan: ' + ', '.join(label for _k, label in PILIHAN[field]) + '.'
            p.write(baris, 0, nama, sel)
            p.write(baris, 1, 'Ya' if harus else '', sel)
            p.write(baris, 2, ket, sel)
        baris += 2
        p.write(baris, 0, 'Daftar bank', tebal)
        p.merge_range(baris, 1, baris, 2, ', '.join(label for _k, label in PILIHAN['bank_tujuan']), sel)
        wb.close()
        return keluaran.getvalue()

    # ------------------------------------------------------------------
    # Pembacaan dan pemeriksaan
    # ------------------------------------------------------------------
    def _baca_berkas(self):
        """[(nomor baris Excel, {field: nilai mentah})] dari lembar data."""
        self.ensure_one()
        if not self.berkas:
            raise UserError(_('Pilih berkas Excel yang akan diimpor.'))
        try:
            import openpyxl  # noqa: PLC0415
            wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(self.berkas)), data_only=True, read_only=True)
        except Exception:
            raise UserError(_('Berkas tidak dapat dibaca. Gunakan berkas Excel (.xlsx) dari templat.'))
        ws = wb[LEMBAR_DATA] if LEMBAR_DATA in wb.sheetnames else wb.worksheets[0]
        baris_iter = ws.iter_rows(values_only=True)
        kepala = next(baris_iter, None) or ()
        kenal = {_kunci(judul): field for field, judul, _w, _j, _k in KOLOM}
        kenal.update({_kunci(field): field for field, _j, _w, _jn, _k in KOLOM})
        posisi = {}
        for i, judul in enumerate(kepala):
            field = kenal.get(_kunci(judul))
            if field and field not in posisi:
                posisi[field] = i
        kurang = [judul for field, judul, harus, _j, _k in KOLOM if harus and field not in posisi]
        if kurang:
            raise UserError(_(
                'Kolom wajib tidak ditemukan di baris pertama: %s. Gunakan templat yang disediakan '
                'dan jangan mengubah judul kolomnya.', ', '.join(kurang)))
        hasil = []
        for nomor, baris in enumerate(baris_iter, start=2):
            nilai = {field: (baris[i] if i < len(baris) else None) for field, i in posisi.items()}
            if all(v is None or str(v).strip() == '' for v in nilai.values()):
                continue
            hasil.append((nomor, nilai))
            if len(hasil) > BARIS_MAKS:
                raise UserError(_('Berkas berisi lebih dari %d baris. Pecah menjadi beberapa berkas.', BARIS_MAKS))
        if not hasil:
            raise UserError(_('Tidak ada data pada lembar "%s". Isi mulai baris ke-2.', ws.title))
        return hasil

    @staticmethod
    def _teks(v):
        """Nilai sel sebagai teks rapi; angka bulat (NIK, NIP yang terbaca sebagai angka) tanpa ',0'."""
        if v is None:
            return ''
        if isinstance(v, float) and v.is_integer():
            v = int(v)
        if isinstance(v, (datetime, date)):
            return v.strftime('%d/%m/%Y')
        return re.sub(r'\s+', ' ', str(v)).strip()

    @staticmethod
    def _angka(v):
        if v is None or v == '':
            return 0.0
        if isinstance(v, (int, float)):
            return float(v)
        teks = re.sub(r'[^\d,.\-]', '', str(v))
        # "8.500.000" atau "8,500,000" -> pemisah ribuan; "8500000,50" -> desimal
        if re.fullmatch(r'-?\d{1,3}([.,]\d{3})+', teks):
            teks = re.sub(r'[.,]', '', teks)
        else:
            teks = teks.replace(',', '.')
        return float(teks)

    @staticmethod
    def _tanggal(v):
        if v is None or v == '':
            return False
        if isinstance(v, datetime):
            return v.date()
        if isinstance(v, date):
            return v
        teks = str(v).strip()
        for pola in ('%d/%m/%Y', '%d-%m-%Y', '%Y-%m-%d', '%d/%m/%y', '%d.%m.%Y'):
            try:
                return datetime.strptime(teks, pola).date()
            except ValueError:
                continue
        raise ValueError(teks)

    def _periksa_baris(self, mentah, nip_berkas, nip_terdaftar):
        """(vals untuk create, [masalah]) dari satu baris."""
        vals, masalah = {}, []
        for field, judul, harus, jenis, _ket in KOLOM:
            v = mentah.get(field)
            try:
                if jenis == 'uang':
                    vals[field] = self._angka(v)
                elif jenis == 'tanggal':
                    vals[field] = self._tanggal(v)
                elif field in ('jenis_kelamin', 'status_pegawai'):
                    teks = self._teks(v)
                    cocok = SINONIM[field].get(_kunci(teks))
                    if teks and not cocok:
                        masalah.append(_('%(kolom)s "%(nilai)s" tidak dikenal', kolom=judul, nilai=teks))
                    vals[field] = cocok or False
                else:
                    vals[field] = self._teks(v) or False
            except (ValueError, TypeError):
                masalah.append(_('%(kolom)s "%(nilai)s" tidak dapat dibaca', kolom=judul, nilai=self._teks(v)))
                vals[field] = False
                continue  # sudah dilaporkan; jangan ditambah "wajib diisi"
            if harus and not vals.get(field):
                masalah.append(_('%s wajib diisi', judul))

        if vals.get('gaji_bersih') and vals['gaji_bersih'] < 0:
            masalah.append(_('Gaji Bersih tidak boleh negatif'))
        if vals.get('simsula') and vals['simsula'] < 0:
            masalah.append(_('Simpanan Sukarela tidak boleh negatif'))
        nip = vals.get('nip')
        if nip:
            if nip_berkas.get(nip.lower(), 0) > 1:
                masalah.append(_('NIP %s muncul lebih dari sekali di berkas', nip))
            if nip.lower() in nip_terdaftar:
                masalah.append(_('NIP %s sudah terdaftar pada mitra ini', nip))
        nik = vals.get('nik')
        if nik:
            nik = re.sub(r'[\s.\-]', '', nik)
            vals['nik'] = nik
            if not re.fullmatch(r'\d{16}', nik):
                masalah.append(_('NIK harus 16 digit angka'))
        telepon = vals.get('telepon')
        if telepon:
            angka = re.sub(r'[^\d+]', '', telepon)
            # Excel sering membuang angka 0 di depan nomor telepon
            if re.fullmatch(r'8\d{8,11}', angka):
                angka = '0' + angka
            vals['telepon'] = angka
        if vals.get('email') and not re.fullmatch(r'[^@\s]+@[^@\s]+\.[^@\s]+', vals['email']):
            masalah.append(_('Email "%s" tidak sah', vals['email']))
        bank = vals.get('bank_tujuan')
        if bank:
            dikenal = {_kunci(label): kode for kode, label in BANK_SELECTION}
            dikenal.update({_kunci(kode): kode for kode, _label in BANK_SELECTION})
            dikenal.update({'bankmandiri': 'mandiri', 'bsi': 'bsi', 'banksyariahindonesia': 'bsi', 'cimb': 'cimb', 'bankbca': 'bca',
                            'bankbri': 'bri', 'bankbni': 'bni', 'bankbtn': 'btn', 'permata': 'permata', 'danamon': 'danamon', 'mega': 'mega'})
            kode = dikenal.get(_kunci(bank))
            if kode and kode != 'lainnya':
                vals['bank_tujuan'] = kode
            else:
                vals['bank_tujuan'], vals['bank_lainnya'] = 'lainnya', bank
        rekening = vals.get('nomor_rekening')
        if rekening and not re.fullmatch(r'[0-9][0-9 \-]{3,29}', rekening):
            masalah.append(_('Nomor Rekening hanya boleh berisi angka'))
        if rekening and not bank:
            masalah.append(_('Bank wajib diisi bila Nomor Rekening diisi'))
        if rekening and not vals.get('atas_nama'):
            vals['atas_nama'] = vals.get('name')
        return {k: v for k, v in vals.items() if v not in (False, None, '', 0.0) or k in ('name', 'nip')}, masalah

    def action_periksa(self):
        self.ensure_one()
        if not self.mitra_id:
            raise UserError(_('Pilih mitra tujuan impor.'))
        if self.mitra_terkunci and self.mitra_id != self.env.user.koperasi_mitra_id:
            raise UserError(_('Anda hanya dapat mengimpor anggota untuk mitra Anda sendiri.'))
        data = self._baca_berkas()
        nip_berkas = {}
        for _nomor, mentah in data:
            nip = self._teks(mentah.get('nip')).lower()
            if nip:
                nip_berkas[nip] = nip_berkas.get(nip, 0) + 1
        # sudo: NIP harus unik per mitra termasuk anggota yang diarsipkan
        nip_terdaftar = {(n or '').lower() for n in self.env['koperasi.anggota'].sudo().with_context(
            active_test=False).search([('mitra_id', '=', self.mitra_id.id)]).mapped('nip')}
        baris = []
        for nomor, mentah in data:
            vals, masalah = self._periksa_baris(mentah, nip_berkas, nip_terdaftar)
            baris.append(fields.Command.create({
                'baris': nomor, 'nama': vals.get('name') or '', 'nip': vals.get('nip') or '',
                'gaji_bersih': vals.get('gaji_bersih') or 0, 'status': 'galat' if masalah else 'ok',
                'pesan': '; '.join(masalah), 'data': json.dumps(vals, default=str),
            }))
        self.write({'state': 'periksa', 'line_ids': [fields.Command.clear()] + baris})
        return self._ulang()

    def action_kembali(self):
        self.write({'state': 'unggah', 'line_ids': [fields.Command.clear()]})
        return self._ulang()

    # ------------------------------------------------------------------
    # Impor
    # ------------------------------------------------------------------
    def action_impor(self):
        self.ensure_one()
        siap = self.line_ids.filtered(lambda l: l.status == 'ok').sorted('baris')
        if not siap:
            raise UserError(_('Tidak ada baris yang siap diimpor. Perbaiki berkasnya lalu unggah ulang.'))
        nilai = []
        for l in siap:
            vals = json.loads(l.data)
            for f in ('tanggal_lahir', 'tanggal_masuk_kerja'):
                if vals.get(f):
                    vals[f] = fields.Date.to_date(vals[f])
            vals['mitra_id'] = self.mitra_id.id
            nilai.append(vals)
        # dengan hak pengguna: admin mitra otomatis terkunci ke mitranya dan status dipaksa Draft
        anggota = self.env['koperasi.anggota'].create(nilai)
        if self.ajukan:
            anggota.action_ajukan()
        self.write({'state': 'selesai', 'jumlah_impor': len(anggota), 'anggota_ids': [fields.Command.set(anggota.ids)]})
        return self._ulang()

    def action_lihat_hasil(self):
        self.ensure_one()
        return {
            'type': 'ir.actions.act_window', 'name': _('Anggota Diimpor'), 'res_model': 'koperasi.anggota',
            'view_mode': 'list,form', 'domain': [('id', 'in', self.anggota_ids.ids)],
        }


class KoperasiImporAnggotaLine(models.TransientModel):
    _name = 'koperasi.impor.anggota.line'
    _description = 'Baris Impor Anggota'
    _order = 'status, baris'

    wizard_id = fields.Many2one('koperasi.impor.anggota.wizard', required=True, ondelete='cascade')
    baris = fields.Integer('Baris Excel')
    nama = fields.Char('Nama')
    nip = fields.Char('NIP')
    gaji_bersih = fields.Float('Gaji Bersih', digits=(16, 0))
    status = fields.Selection([('galat', 'Bermasalah'), ('ok', 'Siap')], string='Hasil')
    pesan = fields.Char('Keterangan')
    data = fields.Text()
